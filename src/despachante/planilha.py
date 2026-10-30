# despachante-digital/src/despachante/planilha.py
"""Planilha Excel consolidada (Etapa 4).

Abas:
  - Processos    : uma linha por processo, com status e valores lado a lado
  - Documentos   : uma linha por PDF extraído
  - Divergencias : uma linha por divergência encontrada
  - Erros        : PDFs que não puderam ser lidos
"""

from __future__ import annotations

from dataclasses import asdict
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from despachante.agrupamento import Processo

VERMELHO = PatternFill("solid", fgColor="F8D7DA")
VERDE = PatternFill("solid", fgColor="D4EDDA")


def _linhas_processos(processos: dict[str, Processo]) -> list[dict]:
    linhas = []
    for proc in processos.values():
        nf, fat = proc.nf, proc.fatura
        linhas.append({
            "Processo": proc.numero,
            "Status": proc.status,
            "Qtd divergencias": len(proc.divergencias),
            "NF": nf.numero if nf else None,
            "Fatura": fat.numero if fat else None,
            "Data NF": nf.data if nf else None,
            "Data fatura": fat.data if fat else None,
            "Valor NF (USD)": nf.valor_total_usd if nf else None,
            "Valor fatura (USD)": fat.valor_total_usd if fat else None,
            "Peso NF (kg)": nf.peso_total_kg if nf else None,
            "Peso fatura (kg)": fat.peso_total_kg if fat else None,
            "NCMs": ", ".join(sorted({n for d in proc.documentos.values() for n in d.ncms})),
        })
    return linhas


def _linhas_documentos(processos: dict[str, Processo]) -> list[dict]:
    linhas = []
    for proc in processos.values():
        for doc in proc.documentos.values():
            d = doc.para_dict()
            d["ncms"] = ", ".join(d["ncms"])
            d["campos_ausentes"] = ", ".join(d["campos_ausentes"])
            linhas.append(d)
    return linhas


def _ajustar_aba(planilha, colorir_status: bool = False) -> None:
    """Cabeçalho em negrito, filtro, painel congelado e largura automática."""
    for celula in planilha[1]:
        celula.font = Font(bold=True)
    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = planilha.dimensions
    for i, coluna in enumerate(planilha.columns, start=1):
        largura = max(len(str(c.value)) if c.value is not None else 0 for c in coluna)
        planilha.column_dimensions[get_column_letter(i)].width = min(largura + 2, 60)
    if colorir_status:
        for linha in planilha.iter_rows(min_row=2):
            cor = VERDE if linha[1].value == "CONFORME" else VERMELHO
            for celula in linha:
                celula.fill = cor


def gerar_planilha(processos: dict[str, Processo], erros: dict[str, dict],
                   destino: Path) -> Path:
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    abas = {
        "Processos": pd.DataFrame(_linhas_processos(processos)),
        "Documentos": pd.DataFrame(_linhas_documentos(processos)),
        "Divergencias": pd.DataFrame(
            [asdict(d) for p in processos.values() for d in p.divergencias],
            columns=["processo", "tipo", "detalhe", "valor_nf", "valor_fatura", "diferenca"]),
        "Erros": pd.DataFrame(
            [{"arquivo": a, **info} for a, info in erros.items()],
            columns=["arquivo", "motivo", "quando"]),
    }
    with pd.ExcelWriter(destino, engine="openpyxl") as escritor:
        for nome, df in abas.items():
            df.to_excel(escritor, sheet_name=nome, index=False)
            _ajustar_aba(escritor.sheets[nome], colorir_status=(nome == "Processos"))
    return destino
